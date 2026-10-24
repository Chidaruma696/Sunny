from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sunny.rules import Condition, Rule, RuleError, Style
from sunny.workbook import apply, formula_for, inspect, matches, output_path, preview, sheet_names


@pytest.fixture
def book(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cobranza"
    ws.append(["Factura", "Cliente", "Estado", "Importe", "Vence"])
    rows = [
        ("F-001", "Acme", "PENDIENTE", 12000, datetime(2026, 1, 15)),
        ("F-002", "Acme", "PAGADO", 80000, datetime(2026, 2, 1)),
        ("F-003", "Globex", "pendiente", 51000, datetime(2026, 3, 10)),
        ("F-004", "Initech", "CANCELADO", None, None),
        (None, "Umbrella", "PAGADO", 49999.5, datetime(2026, 4, 30)),
    ]
    for r in rows:
        ws.append(r)
    ws.append([None, None, None, None, None])  # trailing blank row must not count
    wb.create_sheet("Otra")
    path = tmp_path / "cobranza.xlsx"
    wb.save(path)
    return path


def rule(column, op, value=None, scope="cell", value2=None, values=None, fill="FFC7CE", **style):
    return Rule(column=column, scope=scope, description=f"{column} {op}",
                condition=Condition(op=op, value=value, value2=value2, values=values or []),
                style=Style(fill=fill, **style))


def test_inspect_reads_headers_types_and_rows(book):
    info = inspect(book)
    assert info.name == "Cobranza" and info.first_row == 2 and info.last_row == 6
    kinds = {c.header: c.kind for c in info.columns}
    assert kinds == {"Factura": "text", "Cliente": "text", "Estado": "text", "Importe": "number", "Vence": "date"}
    assert info.column("estado").letter == "C"
    assert info.describe()["rows"] == 5 and "values" not in str(info.describe())
    assert sheet_names(book) == ["Cobranza", "Otra"]
    with pytest.raises(RuleError):
        info.column("Nada")


def test_preview_counts_match_python_semantics(book):
    info = inspect(book)
    counts = {pv.rule.description: pv.cells for pv in preview(book, info, [
        rule("Estado", "eq", "PENDIENTE", scope="row"),
        rule("Importe", "gt", 50000),
        rule("Importe", "between", 10000, value2=60000),
        rule("Factura", "empty"),
        rule("Cliente", "duplicate"),
        rule("Vence", "lt", date(2026, 2, 15)),
        rule("Estado", "in", values=["pagado", "cancelado"]),
        rule("Cliente", "contains", "me"),
    ])}
    assert counts["Estado eq"] == 2          # case-insensitive
    assert counts["Importe gt"] == 2         # 80000 and 51000; None and text never compare
    assert counts["Importe between"] == 3    # 12000, 51000, 49999.5
    assert counts["Factura empty"] == 1
    assert counts["Cliente duplicate"] == 2  # Acme twice
    assert counts["Vence lt"] == 2           # 15 Jan and 1 Feb
    assert counts["Estado in"] == 3
    assert counts["Cliente contains"] == 2   # the two Acme rows; "Umbrella" has no "me"


def test_matches_handles_case_and_kinds():
    r = rule("Estado", "eq", "Pagado")
    assert matches(r, "PAGADO") and not matches(r, "pendiente") and not matches(r, None)
    r.condition.case_sensitive = True
    assert not matches(r, "PAGADO") and matches(r, "Pagado")
    assert not matches(rule("Importe", "gt", 10), "texto")
    assert matches(rule("Importe", "ne", 10), 11) and not matches(rule("Importe", "ne", 10), None)


def test_formulas_pin_the_column_and_guard_numbers(book):
    info = inspect(book)
    col = info.column("Importe")
    assert formula_for(rule("Importe", "gt", 50000), col, 2, 6) == "AND(ISNUMBER($D2),$D2>50000)"
    assert formula_for(rule("Importe", "between", 1, value2=2.5), col, 2, 6) == "AND(ISNUMBER($D2),$D2>=1,$D2<=2.5)"
    estado = info.column("Estado")
    assert formula_for(rule("Estado", "eq", "PENDIENTE"), estado, 2, 6) == '$C2="PENDIENTE"'
    assert formula_for(rule("Estado", "contains", 'a"b'), estado, 2, 6) == 'ISNUMBER(SEARCH("a""b",$C2))'
    assert formula_for(rule("Cliente", "duplicate"), info.column("Cliente"), 2, 6) == 'AND($B2<>"",COUNTIF($B$2:$B$6,$B2)>1)'
    vence = info.column("Vence")
    assert formula_for(rule("Vence", "lt", date(2026, 2, 15)), vence, 2, 6) == "AND(ISNUMBER($E2),$E2<DATE(2026,2,15))"


def test_apply_writes_conditional_formatting_to_a_copy(book):
    info = inspect(book)
    rules = [
        rule("Estado", "eq", "PENDIENTE", scope="row", font_color="9C0006"),
        rule("Importe", "gt", 50000, fill=None, bold=True),
        rule("Cliente", "duplicate", fill="D9D9D9"),
    ]
    rules[2].enabled = False
    out = apply(book, info, rules)
    assert out == output_path(book) and out.name == "cobranza - Sunny.xlsx" and out.exists()
    wb = load_workbook(out)
    ws = wb["Cobranza"]
    ranges = {str(cf.sqref): cf.rules for cf in ws.conditional_formatting}
    assert set(ranges) == {"A2:E6", "D2:D6"}            # disabled rule not written
    row_rule = ranges["A2:E6"][0]
    assert row_rule.formula == ['$C2="PENDIENTE"'] and row_rule.dxf.fill.fgColor.rgb.endswith("FFC7CE")
    bold_rule = ranges["D2:D6"][0]
    assert bold_rule.dxf.font.b is True and bold_rule.dxf.fill is None
    # the original is untouched
    assert not list(load_workbook(book)["Cobranza"].conditional_formatting)
    with pytest.raises(RuleError):
        apply(book, info, rules, book)
